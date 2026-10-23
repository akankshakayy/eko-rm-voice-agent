"""Import Eko's CSP sheets into the knowledge base: the RMs, their CSPs and each CSP's monthly slab.

Usage (the files stay untouched; they are only read):
  python -m app.kb_import data/kb/inbox/<calling sheet>.xlsx              # first import, or a newer sheet
  python -m app.kb_import <file> --remove-demo --map-rm 1=Shiwani          # switch a database off the demo data

What is read from the workbook:
  "Calling Sheet New"  one row per CSP: profile, RM, monthly targets, month-to-date scheme counts and slab
  "Slab"               the scheme counts each slab needs and its monthly incentive

Slab rule (checked against every month in the sheet): a CSP reaches the highest slab where at least 3 of the
4 schemes (PMJDY, PMJJBY, PMSBY, APY) meet that slab's monthly count. `slab_for()` applies it.
Columns are found by their header names, so the sheet can gain months or move columns.
"""
import argparse
import hashlib
import os
import re
import shutil
import sqlite3
import time
from collections import Counter

from . import db

SCHEMES = ("pmjdy", "pmjjby", "pmsby", "apy")
CALLING_TAB, SLAB_TAB = "Calling Sheet New", "Slab"
MONTHS = {m: i for i, m in enumerate(("jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov",
                                      "dec"), 1)}


class ImportError_(Exception):
    pass


# ---------------------------------------------------------------- cleaning cell values
def _text(v) -> str | None:
    if v is None:
        return None
    s = " ".join(str(v).split())
    return s if s and s not in ("-", "—", "–", "NA", "N/A", "#REF!") else None


def _num(v) -> float | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = (_text(v) or "").replace(",", "")
    try:
        return float(s)
    except ValueError:
        return None


def phone10(v) -> str | None:
    """Indian mobile number as 10 digits, or None. Sheets hold them as floats, with spaces or with +91."""
    if isinstance(v, float):
        v = f"{v:.0f}"
    digits = re.sub(r"\D", "", str(v or ""))
    digits = digits[-10:] if len(digits) >= 10 else ""
    return digits if re.fullmatch(r"[6-9]\d{9}", digits) else None


def _email(v) -> str | None:
    s = (_text(v) or "").strip(" ,;").lower()
    return s if re.fullmatch(r"[^@\s,;]+@[^@\s,;]+\.[^@\s,;]+", s) else None


def _slab_name(v) -> str | None:
    """'Slab-1 (₹1,200)', 'Slab 1 ', 'super slab' -> 'Slab-1' / 'Super Slab'; 'No Slab' stays."""
    s = (_text(v) or "").lower()
    if not s:
        return None
    if s.startswith("super"):
        return "Super Slab"
    if s.startswith("no slab"):
        return "No Slab"
    m = re.match(r"slab\s*-?\s*(\d)", s)
    return f"Slab-{m.group(1)}" if m else None


def _category(v) -> str | None:
    """'Diamond' and 'Diamond CSP' are the same tier."""
    s = _text(v)
    tier = (s or "").split(" ")[0].title()
    return f"{tier} CSP" if tier in ("Basic", "Bronze", "Silver", "Gold", "Diamond") else s


def _month(label: str) -> str | None:
    """"July'26", "Aug 2026", "Oct'26 Target" -> '2026-07' etc."""
    m = re.search(r"\b([A-Za-z]{3})[a-z]*\s*'?\s*(\d{4}|\d{2})\b", label or "")
    if not m or m.group(1).lower() not in MONTHS:
        return None
    year = int(m.group(2))
    return f"{year + 2000 if year < 100 else year}-{MONTHS[m.group(1).lower()]:02d}"


# ---------------------------------------------------------------- slab rule
def slab_for(counts: dict, slabs: list[dict]) -> str:
    """Highest slab where at least 3 of the 4 schemes meet the slab's count."""
    got = "No Slab"
    for s in sorted(slabs, key=lambda s: s["rank"]):
        if sum((counts.get(k) or 0) >= s[k] for k in SCHEMES) >= 3:
            got = s["slab"]
    return got


def next_slab_gap(counts: dict, slabs: list[dict]) -> dict | None:
    """The next slab up and the fewest extra counts that reach it (the 3 easiest schemes)."""
    current = slab_for(counts, slabs)
    ranks = {s["slab"]: s["rank"] for s in slabs}
    for s in sorted(slabs, key=lambda s: s["rank"]):
        if s["rank"] <= ranks.get(current, 0):
            continue
        short = {k: max(0, s[k] - int(counts.get(k) or 0)) for k in SCHEMES}
        easiest = dict(sorted(short.items(), key=lambda kv: kv[1])[:3])
        return {"slab": s["slab"], "incentive_rs": s["incentive_rs"],
                "needed": {k: v for k, v in easiest.items() if v > 0}}
    return None


# ---------------------------------------------------------------- reading the workbook
def _rows(ws):
    return [r for r in ws.iter_rows(values_only=True) if any(_text(v) for v in r)]


def read_slab_table(ws) -> list[dict]:
    rows = _rows(ws)
    head = next((i for i, r in enumerate(rows) if any(_slab_name(v) == "Slab-1" for v in r)), None)
    if head is None:
        raise ImportError_(f'Sheet "{SLAB_TAB}": no row with the slab names (Slab-1 ... Super Slab)')
    names = {i: _slab_name(v) for i, v in enumerate(rows[head]) if _slab_name(v)}
    slabs = {n: {"slab": n, "rank": rank} for rank, n in enumerate(names.values(), 1)}
    for r in rows[head + 1:]:
        label = (_text(r[0]) or "").lower()
        key = next((k for k in SCHEMES if label == k), None) or \
            ("incentive_rs" if "incentive" in label else "avg_balance_comm" if "average balance" in label else None)
        if not key:
            continue
        for i, n in names.items():
            v = _num(re.sub(r"[/\-]+$", "", str(r[i] or "")).strip()) if i < len(r) else None
            slabs[n][key] = v if key in ("incentive_rs", "avg_balance_comm") else int(v or 0)
    out = list(slabs.values())
    missing = [k for k in SCHEMES if any(k not in s for s in out)]
    if missing:
        raise ImportError_(f'Sheet "{SLAB_TAB}": no counts for {", ".join(missing).upper()}')
    return out


def _headers(ws):
    """The calling sheet has a group row (e.g. "Targets Oct 2026") above the column names (e.g. "PMJJBY")."""
    top = list(ws.iter_rows(values_only=True, max_row=2))
    if len(top) < 2:
        raise ImportError_(f'Sheet "{CALLING_TAB}" is empty')
    group, cols = None, []
    for i, (g, name) in enumerate(zip(top[0], top[1])):
        if _text(g):
            group = _text(g)
        cols.append((i, group, _text(name) or ""))
    return cols


def _calling_columns(ws) -> tuple[dict, dict]:
    cols = _headers(ws)
    first_group = next((i for i, g, _ in cols if g), len(cols))
    plain = {}
    for i, _, name in cols:
        if i < first_group:
            plain.setdefault(name.lower(), []).append(i)

    def one(*names, last=False):
        for n in names:
            if n.lower() in plain:
                return plain[n.lower()][-1 if last else 0]
        return None

    profile = {"code": one("CSP ID", "CSP Code"), "name": one("CSP Name"), "email": one("CSP Mail ID"),
               "phone": one("CSP Mobile number", "Mobile Number"), "alt_phone": one("Alternative Mobile Number"),
               "category": one("CSP Category"), "rm": one("Relationship Manager"), "rm_email": one("Email of RM"),
               "rm_phone": one("Mobile No Of RM"), "district": one("District"),
               # The sheet has two "State" columns; the first holds the circle head's name, the second the state
               "state": one("State", last=True), "lho": one("Circle (LHO)"), "population": one("Population"),
               "terminal": one("Terminal Status")}
    for k in ("code", "name", "rm"):
        if profile[k] is None:
            raise ImportError_(f'Sheet "{CALLING_TAB}": no column for {k} (header row 2)')

    monthly: dict[str, dict] = {}
    for i, group, name in cols:
        g = group or ""
        first = name.split(" ")[0].lower() if name else ""
        if g.startswith("MTD Achievement") and (m := _month(g)):
            key = first if first in SCHEMES else "slab" if first == "slab" else None
        elif g.startswith("Targets") and (m := _month(g)):
            key = f"target_{first}" if first in SCHEMES else "target_slab" if name.lower().startswith("dfs slab") else None
        elif not group and name.lower().startswith("total transaction count") and (m := _month(name)):
            key = "txn_count"
        elif g == "CSP Commission" and "target" not in name.lower() and (m := _month(name)):
            key = "commission"
        else:
            continue
        if key:
            monthly.setdefault(m, {})[key] = i
    return profile, monthly


def read_workbook(path: str) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)   # read only: the file is never changed
    tabs = {ws.title.strip(): ws for ws in wb.worksheets}
    if CALLING_TAB not in tabs or SLAB_TAB not in tabs:
        raise ImportError_(f'Expected the sheets "{CALLING_TAB}" and "{SLAB_TAB}"; found {", ".join(tabs)}')
    slabs = read_slab_table(tabs[SLAB_TAB])
    ws = tabs[CALLING_TAB]
    profile, monthly = _calling_columns(ws)
    get = lambda r, k: r[profile[k]] if profile.get(k) is not None and profile[k] < len(r) else None
    csps, seen = [], set()
    for r in ws.iter_rows(values_only=True, min_row=3):
        code = (_text(get(r, "code")) or "").upper()
        if not re.fullmatch(r"[0-9A-Z]{6,12}", code) or code in seen:
            continue
        seen.add(code)
        months = {}
        for m, idx in monthly.items():
            vals = {k: (_slab_name(r[i]) if k in ("slab", "target_slab") else _num(r[i]))
                    for k, i in idx.items() if i < len(r)}
            if any(v is not None for v in vals.values()):
                months[m] = vals
        csps.append({
            "csp_code": code, "name": _text(get(r, "name")), "email": _email(get(r, "email")),
            "phone": phone10(get(r, "phone")), "alt_phone": phone10(get(r, "alt_phone")),
            "category": _category(get(r, "category")), "rm": (_text(get(r, "rm")) or "").title() or None,
            "rm_email": _email(get(r, "rm_email")), "rm_phone": phone10(get(r, "rm_phone")),
            "district": _text(get(r, "district")), "state": (_text(get(r, "state")) or "").title() or None,
            "lho": _text(get(r, "lho")), "population": (_text(get(r, "population")) or "").title() or None,
            "terminal_status": _text(get(r, "terminal")), "months": months})
    return {"slabs": slabs, "csps": csps, "months": sorted(monthly)}


# ---------------------------------------------------------------- checks
def check(data: dict) -> list[str]:
    """Things worth knowing about a sheet before trusting it; stored with the import."""
    notes, csps = [], data["csps"]
    phones = Counter(c["phone"] for c in csps if c["phone"])
    shared = sum(1 for c in csps if c["phone"] and phones[c["phone"]] > 1)
    if shared:
        notes.append(f"{shared} CSPs share a mobile number with another CSP; those numbers won't identify callers")
    if n := sum(1 for c in csps if not c["phone"]):
        notes.append(f"{n} CSPs have no valid mobile number")
    if n := sum(1 for c in csps if not c["rm"]):
        notes.append(f"{n} CSPs have no RM")
    for m in data["months"]:
        rows = [c["months"][m] for c in csps if m in c["months"] and c["months"][m].get("slab")]
        bad = sum(1 for v in rows if slab_for(v, data["slabs"]) != v["slab"])
        if rows:
            notes.append(f"{m}: slab column {'matches' if not bad else f'differs from'} the 3-of-4 rule"
                         + (f" for {bad} of {len(rows)} CSPs" if bad else f" for all {len(rows)} CSPs"))
    return notes


# ---------------------------------------------------------------- writing
def _sha256(path: str) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def _demo_counts() -> dict:
    return {t: db.q(f"SELECT COUNT(*) n FROM {t} WHERE source='demo'", one=True)["n"] for t in ("rms", "csps")}


def import_file(path: str, by: str = "cli", force: bool = False) -> dict:
    sha = _sha256(path)
    old = db.q("SELECT * FROM kb_imports WHERE sha256=?", (sha,), one=True)
    if old and not force:
        return {"skipped": True, "reason": f"already imported (import {old['id']})", "import_id": old["id"]}
    data = read_workbook(path)
    if _sha256(path) != sha:
        raise ImportError_("The file changed while it was being read; try again")
    if not data["csps"]:
        raise ImportError_("No CSP rows found")
    notes = check(data)
    now = time.time()
    phones = Counter(c["phone"] for c in data["csps"] if c["phone"])
    c = db.conn()
    with db._lock, c:                                   # one transaction: all of it or nothing
        cur = c.execute("INSERT OR REPLACE INTO kb_imports(file_name,sha256,kind,rows,notes,imported_by,imported_at)"
                        " VALUES(?,?,?,?,?,?,?)", (os.path.basename(path), sha, "calling_sheet", len(data["csps"]),
                                                     "\n".join(notes), by, now))
        imp = cur.lastrowid
        for s in data["slabs"]:
            c.execute("INSERT OR REPLACE INTO slabs(slab,rank,pmjdy,pmjjby,pmsby,apy,incentive_rs,avg_balance_comm,"
                      "import_id) VALUES(?,?,?,?,?,?,?,?,?)", (s["slab"], s["rank"], s["pmjdy"], s["pmjjby"],
                                                              s["pmsby"], s["apy"], s.get("incentive_rs"),
                                                              s.get("avg_balance_comm"), imp))
        # RMs: matched by name; email and phone are the ones most of their CSP rows give
        rm_ids = {}
        for name in sorted({x["rm"] for x in data["csps"] if x["rm"]}):
            mine = [x for x in data["csps"] if x["rm"] == name]
            email = (Counter(x["rm_email"] for x in mine if x["rm_email"]).most_common(1) or [(None,)])[0][0]
            phone = (Counter(x["rm_phone"] for x in mine if x["rm_phone"]).most_common(1) or [(None,)])[0][0]
            region = ", ".join(k for k, _ in Counter(x["lho"] for x in mine if x["lho"]).most_common(2))
            row = c.execute("SELECT id FROM rms WHERE source='kb' AND lower(name)=lower(?)", (name,)).fetchone()
            if row:
                c.execute("UPDATE rms SET email=coalesce(?,email), phone=coalesce(?,phone), region=? WHERE id=?",
                          (email, phone, region, row[0]))
                rm_ids[name] = row[0]
            else:
                rm_ids[name] = c.execute("INSERT INTO rms(name,phone,region,email,source) VALUES(?,?,?,?,'kb')",
                                         (name, phone, region, email)).lastrowid
        # Numbers are re-assigned from this sheet; clear them first so a number that moved can't collide
        c.execute("UPDATE csps SET phone=NULL WHERE source='kb'")
        for x in data["csps"]:
            own = x["phone"] if x["phone"] and phones[x["phone"]] == 1 else None
            shared = x["phone"] if x["phone"] and phones[x["phone"]] > 1 else None
            if own:   # a demo CSP can't hold a real number; demo data is removed before real use anyway
                c.execute("UPDATE csps SET phone=NULL WHERE phone=? AND source='demo'", (own,))
            vals = dict(name=x["name"], phone=own, shared_phone=shared, alt_phone=x["alt_phone"], email=x["email"],
                        category=x["category"], rm_id=rm_ids.get(x["rm"]), district=x["district"], state=x["state"],
                        lho=x["lho"], population=x["population"], terminal_status=x["terminal_status"], active=1,
                        source="kb", import_id=imp, updated_at=now)
            row = c.execute("SELECT id FROM csps WHERE csp_code=?", (x["csp_code"],)).fetchone()
            if row:
                c.execute(f"UPDATE csps SET {', '.join(f'{k}=?' for k in vals)} WHERE id=?", (*vals.values(), row[0]))
                csp_id = row[0]
            else:
                csp_id = c.execute(f"INSERT INTO csps(csp_code,{','.join(vals)}) VALUES(?{',?' * len(vals)})",
                                   (x["csp_code"], *vals.values())).lastrowid
            for month, v in x["months"].items():
                cols = ["csp_id", "month", *v, "import_id", "updated_at"]
                c.execute(f"INSERT INTO csp_monthly({','.join(cols)}) VALUES({','.join('?' * len(cols))}) "
                          f"ON CONFLICT(csp_id,month) DO UPDATE SET {', '.join(f'{k}=excluded.{k}' for k in cols[2:])}",
                          (csp_id, month, *v.values(), imp, now))
        # CSPs that left the sheet stay (their tickets and calls point at them) but are marked inactive
        codes = [x["csp_code"] for x in data["csps"]]
        gone = c.execute(f"UPDATE csps SET active=0, updated_at=? WHERE source='kb' AND active=1 AND csp_code NOT IN "
                         f"({','.join('?' * len(codes))})", (now, *codes)).rowcount
    return {"import_id": imp, "file": os.path.basename(path), "sha256": sha, "csps": len(data["csps"]),
            "rms": {n: sum(1 for x in data["csps"] if x["rm"] == n) for n in rm_ids}, "months": data["months"],
            "now_inactive": gone, "notes": notes}


# ---------------------------------------------------------------- leaving the demo data behind
def remove_demo(rm_map: dict[int, int] | None = None) -> dict:
    """Move a database off the made-up RMs/CSPs: copy it to demo.db (for the simulator and tests), then delete
    the demo RMs, CSPs and every simulated call and ticket here. Logins, upload tokens and recordings linked to a
    demo RM move to the real RM in `rm_map` ({demo_id: real_id}), otherwise they are unlinked."""
    rm_map = rm_map or {}
    demo_rms = [r["id"] for r in db.q("SELECT id FROM rms WHERE source='demo'")]
    demo_csps = [r["id"] for r in db.q("SELECT id FROM csps WHERE source='demo'")]
    if not demo_rms and not demo_csps:
        return {"removed": False}
    for real in rm_map.values():
        if not db.q("SELECT id FROM rms WHERE id=? AND source='kb'", (real,), one=True):
            raise ImportError_(f"RM {real} is not a real (imported) RM")
    backup = _demo_copy()
    ids = lambda xs: ",".join(str(int(i)) for i in xs) or "NULL"
    c = db.conn()
    out = {"removed": True, "demo_copy": backup}
    with db._lock, c:
        calls = [r[0] for r in c.execute(f"SELECT id FROM calls WHERE simulated=1 OR csp_id IN ({ids(demo_csps)}) "
                                         f"OR rm_id IN ({ids(demo_rms)})")]
        tickets = [r[0] for r in c.execute(f"SELECT id FROM tickets WHERE call_id IN ({ids(calls)}) "
                                           f"OR csp_id IN ({ids(demo_csps)}) OR rm_id IN ({ids(demo_rms)})")]
        c.execute(f"DELETE FROM ticket_events WHERE ticket_id IN ({ids(tickets)})")
        c.execute(f"DELETE FROM tickets WHERE id IN ({ids(tickets)})")
        c.execute(f"DELETE FROM turns WHERE call_id IN ({ids(calls)})")
        c.execute(f"DELETE FROM calls WHERE id IN ({ids(calls)})")
        out.update(calls_deleted=len(calls), tickets_deleted=len(tickets))
        c.execute(f"UPDATE recordings SET csp_id=NULL WHERE csp_id IN ({ids(demo_csps)})")
        moved, unlinked = [], []
        for rid in demo_rms:
            real = rm_map.get(rid)
            for table in ("users", "upload_tokens", "recordings"):
                n = c.execute(f"UPDATE {table} SET rm_id=? WHERE rm_id=?", (real, rid)).rowcount
                (moved if real else unlinked).extend([f"{table}:{rid}"] * n)
        # An RM login whose demo profile is gone waits for an admin to link it to a real RM
        c.execute("UPDATE users SET role='pending' WHERE role='rm' AND rm_id IS NULL")
        n_tok = c.execute("UPDATE upload_tokens SET revoked_at=? WHERE rm_id IS NULL AND revoked_at IS NULL",
                          (time.time(),)).rowcount
        c.execute(f"DELETE FROM csps WHERE id IN ({ids(demo_csps)})")
        c.execute(f"DELETE FROM rms WHERE id IN ({ids(demo_rms)})")
        out.update(rms_deleted=len(demo_rms), csps_deleted=len(demo_csps), moved=Counter(moved),
                   unlinked=Counter(unlinked), upload_tokens_revoked=n_tok)
    return out


def _demo_copy() -> str:
    """Copy the database (made-up data only) to demo.db next to it; real recordings and secrets are left out."""
    src = os.path.abspath(db.DB_PATH)
    dest = os.path.join(os.path.dirname(src), "demo.db")
    if os.path.exists(dest):
        dest = os.path.join(os.path.dirname(src), f"demo-{time.strftime('%Y%m%d-%H%M%S')}.db")
    with db._lock:
        out = sqlite3.connect(dest)
        db.conn().backup(out)
        for t in ("recordings", "upload_tokens", "sessions", "auth_tokens", "api_keys", "webhooks",
                  "webhook_deliveries", "kb_imports", "csp_monthly", "slabs"):
            out.execute(f"DELETE FROM {t}")
        out.execute("DELETE FROM csps WHERE source<>'demo'")
        out.execute("DELETE FROM rms WHERE source<>'demo'")
        out.commit()
        out.execute("VACUUM")   # so deleted real rows don't linger in free pages of the copy
        out.close()
    return dest


def main():
    p = argparse.ArgumentParser(description="Import Eko CSP sheets into Sahayak's knowledge base")
    p.add_argument("files", nargs="+")
    p.add_argument("--force", action="store_true", help="import again even if this exact file was imported")
    p.add_argument("--remove-demo", action="store_true", help="delete the made-up RMs/CSPs and simulated calls")
    p.add_argument("--map-rm", action="append", default=[], metavar="DEMO_ID=NAME",
                   help="move logins/upload tokens of a demo RM to a real RM, e.g. 1=Shiwani")
    a = p.parse_args()
    demo = _demo_counts()
    if any(demo.values()) and not a.remove_demo:
        raise SystemExit(f"This database still has demo data ({demo['rms']} RMs, {demo['csps']} CSPs). Real and "
                         "demo data must not mix: run again with --remove-demo (a copy is kept as demo.db).")
    for f in a.files:
        r = import_file(f, force=a.force)
        print(f"{f}: " + (r["reason"] if r.get("skipped") else
                          f"{r['csps']} CSPs, RMs {r['rms']}, months {', '.join(r['months'])}, "
                          f"{r['now_inactive']} marked inactive"))
        for n in r.get("notes", []):
            print("  -", n)
    if a.remove_demo:
        rm_map = {}
        for m in a.map_rm:
            demo_id, _, name = m.partition("=")
            row = db.q("SELECT id FROM rms WHERE source='kb' AND lower(name)=lower(?)", (name.strip(),), one=True)
            if not row:
                raise SystemExit(f"No imported RM called {name!r}")
            rm_map[int(demo_id)] = row["id"]
        print(remove_demo(rm_map))


if __name__ == "__main__":
    main()
